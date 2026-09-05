import sqlite3
from contextlib import closing
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import reporting


class ReportingTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.db = Path(self.directory.name) / 'attempts.db'
        self.report = Path(self.directory.name) / 'report.xlsx'
        self.problem = dict(frontendQuestionId='1512', title='Good Pairs',
                            titleSlug='number-of-good-pairs', difficulty='Easy')

    def rows(self, table):
        with closing(sqlite3.connect(self.db)) as conn:
            conn.row_factory = sqlite3.Row
            return [dict(row) for row in conn.execute(f'SELECT * FROM {table}')]

    def test_start_session_does_not_interrupt_other_records(self):
        first = reporting.start_session(self.db, ('easy',), 'human', 3)
        reporting.start_attempt(self.db, first, self.problem)
        reporting.start_session(self.db, ('hard',), 'instant', 1)
        self.assertEqual([row['status'] for row in self.rows('sessions')], ['Running', 'Running'])
        self.assertEqual(self.rows('attempts')[0]['outcome'], 'In progress')

    def test_finish_session_closes_only_its_unfinished_attempts(self):
        session = reporting.start_session(self.db, ('easy',), 'instant', 1)
        accepted, started = reporting.start_attempt(self.db, session, self.problem)
        reporting.finish_attempt(self.db, accepted, started, 'Submission', 'Accepted')
        reporting.start_attempt(self.db, session, self.problem)
        other = reporting.start_session(self.db, ('easy',), 'instant', 1)
        reporting.start_attempt(self.db, other, self.problem)
        reporting.finish_session(self.db, session, 'Interrupted', 'Stopped by user')
        self.assertEqual([row['outcome'] for row in self.rows('attempts')],
                         ['Accepted', 'Interrupted', 'In progress'])

    def test_explicit_recovery_preserves_finished_records_and_last_stage(self):
        first = reporting.start_session(self.db, ('easy',), 'human', 3)
        attempt, _ = reporting.start_attempt(self.db, first, self.problem)
        reporting.update_attempt_stage(self.db, attempt, 'Submission')
        reporting.record_submission(self.db, attempt, 123)
        completed = reporting.start_session(self.db, ('easy',), 'instant', 1)
        reporting.finish_session(self.db, completed, 'Completed')
        reporting.recover_interrupted_sessions(self.db)
        attempt = self.rows('attempts')[0]
        self.assertEqual(attempt['outcome'], 'Interrupted')
        self.assertEqual(attempt['stage'], 'Submission')
        self.assertEqual(attempt['submission_id'], '123')
        self.assertEqual([row['status'] for row in self.rows('sessions')], ['Interrupted', 'Completed'])

    def test_poll_failure_keeps_submission_id_and_export_counts_network_error(self):
        session = reporting.start_session(self.db, ('easy',), 'instant', 1)
        attempt, start = reporting.start_attempt(self.db, session, self.problem)
        reporting.record_submission(self.db, attempt, 321)
        reporting.finish_attempt(self.db, attempt, start, 'Submission', 'Network error', reason='offline')
        reporting.finish_session(self.db, session, 'Network error', 'offline')
        self.assertEqual(self.rows('attempts')[0]['submission_id'], '321')
        reporting.export_excel(self.db, self.report)
        wb = load_workbook(self.report)
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames, ['Summary', 'Attempts', 'Sessions', 'Historical solved IDs'])
        headers = [cell.value for cell in wb['Sessions'][1]]
        row = dict(zip(headers, [cell.value for cell in wb['Sessions'][2]]))
        self.assertEqual(row['Network errors'], 1)
        self.assertEqual(row['Status'], 'Network error')
        self.assertEqual(wb['Attempts']['L2'].value, '321')
        self.assertIsNotNone(wb['Attempts']['N2'].hyperlink)
        self.assertEqual(wb['Sessions'].freeze_panes, 'A2')

    def test_failed_export_does_not_replace_previous_workbook_or_history(self):
        session = reporting.start_session(self.db, ('easy',), 'instant', 1)
        reporting.finish_session(self.db, session, 'Completed')
        reporting.export_excel(self.db, self.report)
        previous = self.report.read_bytes()
        with patch.object(reporting.os, 'replace', side_effect=PermissionError('open workbook')):
            with self.assertRaises(PermissionError):
                reporting.export_excel(self.db, self.report)
        self.assertEqual(self.report.read_bytes(), previous)
        self.assertEqual(self.rows('sessions')[0]['status'], 'Completed')

    def test_historical_import_does_not_duplicate_modern_acceptances(self):
        session = reporting.start_session(self.db, ('easy',), 'instant', 1)
        attempt, start = reporting.start_attempt(self.db, session, self.problem)
        reporting.finish_attempt(self.db, attempt, start, 'Submission', 'Accepted')
        reporting.import_historical_solved(self.db, ['1512', '1', '1'])
        self.assertEqual([row['problem_id'] for row in self.rows('historical_solved')], ['1'])
        self.assertEqual(reporting.accepted_problem_ids(self.db), ['1512'])
