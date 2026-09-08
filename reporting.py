"""Durable run history and Excel reporting for the LeetCode bot."""

import os
import re
import sqlite3
import time
import uuid
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path


OUTCOME_COLORS = {
    "Accepted": ("C6EFCE", "006100"),
    "No solution": ("DDEBF7", "1F4E78"),
    "Test failed": ("FFF2CC", "7F6000"),
    "Submission failed": ("FCE4D6", "9C5700"),
    "Rate limited": ("E4DFEC", "5F497A"),
    "Browser error": ("F4CCCC", "9C0006"),
    "Network error": ("F4CCCC", "9C0006"),
    "Error": ("F4CCCC", "9C0006"),
    "Interrupted": ("E7E6E6", "595959"),
    "In progress": ("E7E6E6", "595959"),
}

SESSION_COLORS = {
    "Completed": ("C6EFCE", "006100"),
    "Partial": ("FFF2CC", "7F6000"),
    "Rate limited": ("E4DFEC", "5F497A"),
    "Browser error": ("F4CCCC", "9C0006"),
    "Network error": ("F4CCCC", "9C0006"),
    "Failed": ("F4CCCC", "9C0006"),
    "Interrupted": ("E7E6E6", "595959"),
    "Running": ("E7E6E6", "595959"),
}


def now_iso():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def clean_reason(value, limit=500):
    """Keep report reasons useful without storing entire HTML responses."""
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if len(text) > limit:
        return text[: limit - 1] + "…"
    return text


@contextmanager
def _connect(db_path):
    db_path = Path(db_path)
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(
        """
        PRAGMA journal_mode = WAL;

        CREATE TABLE IF NOT EXISTS sessions (
            session_id TEXT PRIMARY KEY,
            started_at TEXT NOT NULL,
            ended_at TEXT,
            difficulties TEXT NOT NULL,
            timing TEXT NOT NULL,
            target_count INTEGER,
            status TEXT NOT NULL DEFAULT 'Running',
            stop_reason TEXT NOT NULL DEFAULT ''
        );

        CREATE TABLE IF NOT EXISTS attempts (
            attempt_id INTEGER PRIMARY KEY AUTOINCREMENT,
            session_id TEXT NOT NULL,
            started_at TEXT NOT NULL,
            ended_at TEXT,
            problem_id TEXT NOT NULL,
            title TEXT NOT NULL,
            slug TEXT NOT NULL,
            difficulty TEXT NOT NULL,
            stage TEXT NOT NULL DEFAULT 'Selected',
            outcome TEXT NOT NULL DEFAULT 'In progress',
            reason TEXT NOT NULL DEFAULT '',
            status_message TEXT NOT NULL DEFAULT '',
            runtime TEXT NOT NULL DEFAULT '',
            submission_id TEXT NOT NULL DEFAULT '',
            duration_seconds REAL,
            problem_url TEXT NOT NULL,
            solution_url TEXT NOT NULL,
            FOREIGN KEY (session_id) REFERENCES sessions(session_id)
        );

        CREATE TABLE IF NOT EXISTS historical_solved (
            problem_id TEXT PRIMARY KEY,
            imported_at TEXT NOT NULL,
            source TEXT NOT NULL
        );

        CREATE INDEX IF NOT EXISTS attempts_session_idx
            ON attempts(session_id);
        CREATE INDEX IF NOT EXISTS attempts_outcome_idx
            ON attempts(outcome);
        """
    )
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def initialize(db_path):
    with _connect(db_path):
        pass


def import_historical_solved(db_path, problem_ids):
    """Preserve old solved IDs without inventing dates, titles, or outcomes."""
    rows = [(str(problem_id), now_iso(), "solved.json") for problem_id in problem_ids]
    if not rows:
        initialize(db_path)
        return
    with _connect(db_path) as conn:
        conn.executemany(
            """
            INSERT OR IGNORE INTO historical_solved(problem_id, imported_at, source)
            SELECT ?, ?, ?
            WHERE NOT EXISTS (
                SELECT 1 FROM attempts WHERE problem_id = ? AND outcome = 'Accepted'
            )
            """,
            [(*row, row[0]) for row in rows],
        )


def recover_interrupted_sessions(db_path):
    """Recover abandoned records only while the caller holds the run lock."""
    with _connect(db_path) as conn:
        recovered_at = now_iso()
        recovery_reason = "Previous process ended before this record was completed"
        conn.execute(
            """
            UPDATE attempts
            SET ended_at = ?, outcome = 'Interrupted', reason = ?,
                duration_seconds = NULL
            WHERE outcome = 'In progress'
            """,
            (recovered_at, recovery_reason),
        )
        conn.execute(
            """
            UPDATE sessions
            SET ended_at = ?, status = 'Interrupted', stop_reason = ?
            WHERE status = 'Running'
            """,
            (recovered_at, recovery_reason),
        )


def start_session(db_path, difficulties, timing, target_count):
    session_id = f"{datetime.now():%Y%m%d-%H%M%S}-{uuid.uuid4().hex[:6]}"
    with _connect(db_path) as conn:
        conn.execute(
            """
            INSERT INTO sessions(
                session_id, started_at, difficulties, timing, target_count
            ) VALUES (?, ?, ?, ?, ?)
            """,
            (
                session_id,
                now_iso(),
                ",".join(difficulties),
                timing,
                target_count,
            ),
        )
    return session_id


def accepted_problem_ids(db_path):
    with _connect(db_path) as conn:
        return [row[0] for row in conn.execute(
            "SELECT DISTINCT problem_id FROM attempts WHERE outcome = 'Accepted'"
        )]


def last_session_at(db_path):
    """Most recent session start, or None when nothing has run yet."""
    if not Path(db_path).exists():
        return None
    with _connect(db_path) as conn:
        row = conn.execute("SELECT MAX(started_at) FROM sessions").fetchone()
    return row[0] if row and row[0] else None


def update_attempt_stage(db_path, attempt_id, stage):
    with _connect(db_path) as conn:
        conn.execute("UPDATE attempts SET stage = ? WHERE attempt_id = ?",
                     (stage, attempt_id))


def record_submission(db_path, attempt_id, submission_id):
    """Keep the ID even if polling is interrupted or the network disappears."""
    with _connect(db_path) as conn:
        conn.execute("UPDATE attempts SET submission_id = ? WHERE attempt_id = ?",
                     (str(submission_id), attempt_id))


def finish_session(db_path, session_id, status, stop_reason=""):
    with _connect(db_path) as conn:
        # Also cover interruption immediately after an attempt INSERT, before
        # the caller received its ID. Never rewrite an already finished result.
        outcome = status if status in OUTCOME_COLORS else "Error"
        conn.execute(
            """
            UPDATE attempts SET ended_at = ?, outcome = ?, reason = ?
            WHERE session_id = ? AND outcome = 'In progress'
            """,
            (now_iso(), outcome,
             clean_reason(stop_reason or "Session ended before attempt was finalized"),
             session_id),
        )
        conn.execute(
            """
            UPDATE sessions
            SET ended_at = ?, status = ?, stop_reason = ?
            WHERE session_id = ?
            """,
            (now_iso(), status, clean_reason(stop_reason), session_id),
        )


def start_attempt(db_path, session_id, problem):
    problem_id = str(problem["frontendQuestionId"])
    slug = problem["titleSlug"]
    with _connect(db_path) as conn:
        cursor = conn.execute(
            """
            INSERT INTO attempts(
                session_id, started_at, problem_id, title, slug, difficulty,
                problem_url, solution_url
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                session_id,
                now_iso(),
                problem_id,
                problem["title"],
                slug,
                problem["difficulty"].title(),
                f"https://leetcode.com/problems/{slug}/",
                f"https://walkccc.me/LeetCode/problems/{problem_id}/",
            ),
        )
        return cursor.lastrowid, time.monotonic()


def finish_attempt(
    db_path,
    attempt_id,
    started_monotonic,
    stage,
    outcome,
    reason="",
    status_message="",
    runtime="",
    submission_id="",
):
    duration = max(0.0, time.monotonic() - started_monotonic)
    with _connect(db_path) as conn:
        conn.execute(
            """
            UPDATE attempts
            SET ended_at = ?, stage = ?, outcome = ?, reason = ?,
                status_message = ?, runtime = ?,
                submission_id = CASE WHEN ? = '' THEN submission_id ELSE ? END,
                duration_seconds = ?
            WHERE attempt_id = ?
            """,
            (
                now_iso(),
                stage,
                outcome,
                clean_reason(reason),
                clean_reason(status_message, limit=200),
                str(runtime or ""),
                str(submission_id or ""),
                str(submission_id or ""),
                round(duration, 3),
                attempt_id,
            ),
        )


def _excel_datetime(value):
    if not value:
        return None
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone().replace(tzinfo=None)
    return parsed


def _style_sheet(ws, widths, date_columns=()):
    from openpyxl.styles import Alignment, Font, PatternFill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="222222")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for column in date_columns:
        for cell in ws[column][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"


def _add_table(ws, name):
    from openpyxl.worksheet.table import Table, TableStyleInfo

    if ws.max_row < 2:
        return
    table = Table(displayName=name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _add_status_formatting(ws, status_column, end_column, colors):
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill

    last_row = max(2, ws.max_row)
    target = f"A2:{end_column}{last_row}"
    for status, (fill_color, font_color) in colors.items():
        ws.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'${status_column}2="{status}"'],
                fill=PatternFill("solid", fgColor=fill_color),
                font=Font(color=font_color),
            ),
        )


def export_excel(db_path, report_path, account_name=None):
    """Regenerate a styled Excel report from the SQLite source of truth."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "Excel export requires openpyxl; run `python -m pip install -r requirements.txt`"
        ) from exc

    initialize(db_path)
    with _connect(db_path) as conn:
        attempts = conn.execute(
            "SELECT * FROM attempts ORDER BY attempt_id"
        ).fetchall()
        sessions = conn.execute(
            """
            SELECT s.*,
                COALESCE(SUM(a.outcome = 'Accepted'), 0) AS accepted,
                COALESCE(SUM(a.outcome = 'No solution'), 0) AS no_solution,
                COALESCE(SUM(a.outcome = 'Test failed'), 0) AS test_failed,
                COALESCE(SUM(a.outcome = 'Submission failed'), 0) AS submit_failed,
                COALESCE(SUM(a.outcome = 'Rate limited'), 0) AS rate_limited,
                COALESCE(SUM(a.outcome = 'Browser error'), 0) AS browser_error,
                COALESCE(SUM(a.outcome = 'Network error'), 0) AS network_error,
                COALESCE(SUM(a.outcome = 'Error'), 0) AS errors,
                COALESCE(SUM(a.outcome = 'Interrupted'), 0) AS interrupted
            FROM sessions s
            LEFT JOIN attempts a ON a.session_id = s.session_id
            GROUP BY s.session_id
            ORDER BY s.started_at DESC, s.rowid DESC
            """
        ).fetchall()
        historical = conn.execute(
            "SELECT * FROM historical_solved ORDER BY CAST(problem_id AS INTEGER)"
        ).fetchall()
        outcome_counts = {
            row["outcome"]: row["count"]
            for row in conn.execute(
                "SELECT outcome, COUNT(*) AS count FROM attempts GROUP BY outcome"
            )
        }

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    if account_name:
        wb.properties.title = f"LeetCode report — {account_name}"
    summary.sheet_view.showGridLines = False
    summary["A1"] = f"LeetCode bot report — {account_name}" if account_name else "LeetCode bot report"
    summary["A1"].font = Font(name="Arial", size=16, bold=True, color="1F1F1F")
    summary["A2"] = f"Generated {datetime.now():%Y-%m-%d %H:%M:%S}"
    summary["A2"].font = Font(name="Arial", size=10, italic=True, color="666666")
    summary["A4"] = "Outcome"
    summary["B4"] = "Attempts"
    for cell in summary[4]:
        if cell.column <= 2:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")
    summary_outcomes = list(OUTCOME_COLORS)
    for row_number, outcome in enumerate(summary_outcomes, start=5):
        summary.cell(row_number, 1, outcome)
        summary.cell(row_number, 2, outcome_counts.get(outcome, 0))
        fill_color, font_color = OUTCOME_COLORS[outcome]
        summary.cell(row_number, 1).fill = PatternFill("solid", fgColor=fill_color)
        summary.cell(row_number, 1).font = Font(name="Arial", color=font_color)
        summary.cell(row_number, 2).number_format = "#,##0"
    summary["D4"] = "Other records"
    summary["E4"] = "Count"
    for cell in summary[4][3:5]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for row_number, (label, value) in enumerate(
        (
            ("Sessions", len(sessions)),
            ("Total attempts", len(attempts)),
            ("Historical solved IDs", len(historical)),
        ),
        start=5,
    ):
        summary.cell(row_number, 4, label)
        summary.cell(row_number, 5, value).number_format = "#,##0"
    for row in summary.iter_rows(min_row=5, max_row=4 + len(summary_outcomes), min_col=1, max_col=5):
        for cell in row:
            if cell.value is not None and cell.column != 1:
                cell.font = Font(name="Arial", size=10, color="222222")
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 14
    summary.column_dimensions["C"].width = 3
    summary.column_dimensions["D"].width = 24
    summary.column_dimensions["E"].width = 12

    attempts_ws = wb.create_sheet("Attempts")
    attempt_headers = [
        "Started", "Ended", "Session ID", "Problem ID", "Title", "Difficulty",
        "Stage", "Outcome", "Reason", "LeetCode status", "Runtime",
        "Submission ID", "Duration (s)", "Problem URL", "Solution URL",
    ]
    attempts_ws.append(attempt_headers)
    for row in attempts:
        attempts_ws.append(
            [
                _excel_datetime(row["started_at"]),
                _excel_datetime(row["ended_at"]),
                row["session_id"],
                row["problem_id"],
                row["title"],
                row["difficulty"],
                row["stage"],
                row["outcome"],
                row["reason"],
                row["status_message"],
                row["runtime"],
                row["submission_id"],
                row["duration_seconds"],
                row["problem_url"],
                row["solution_url"],
            ]
        )
    _style_sheet(
        attempts_ws,
        {
            "A": 20, "B": 20, "C": 25, "D": 11, "E": 34, "F": 11,
            "G": 18, "H": 19, "I": 55, "J": 22, "K": 12, "L": 17,
            "M": 13, "N": 48, "O": 48,
        },
        date_columns=("A", "B"),
    )
    for cell in attempts_ws["M"][1:]:
        cell.number_format = "0.0"
    for row_number in range(2, attempts_ws.max_row + 1):
        for column in ("N", "O"):
            cell = attempts_ws[f"{column}{row_number}"]
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(
                    name="Arial", size=10, color="0563C1", underline="single"
                )
    _add_status_formatting(attempts_ws, "H", "O", OUTCOME_COLORS)
    _add_table(attempts_ws, "AttemptsTable")

    sessions_ws = wb.create_sheet("Sessions")
    session_headers = [
        "Session ID", "Started", "Ended", "Difficulties", "Timing", "Target",
        "Accepted", "No solution", "Test failed", "Submission failed",
        "Rate limited", "Browser errors", "Network errors", "Errors", "Interrupted", "Status",
        "Stop reason",
    ]
    sessions_ws.append(session_headers)
    for row in sessions:
        sessions_ws.append(
            [
                row["session_id"],
                _excel_datetime(row["started_at"]),
                _excel_datetime(row["ended_at"]),
                row["difficulties"],
                row["timing"],
                row["target_count"],
                row["accepted"],
                row["no_solution"],
                row["test_failed"],
                row["submit_failed"],
                row["rate_limited"],
                row["browser_error"],
                row["network_error"],
                row["errors"],
                row["interrupted"],
                row["status"],
                row["stop_reason"],
            ]
        )
    _style_sheet(
        sessions_ws,
        {
            "A": 25, "B": 20, "C": 20, "D": 18, "E": 12, "F": 10,
            "G": 11, "H": 13, "I": 12, "J": 18, "K": 14, "L": 15,
            "M": 15, "N": 10, "O": 12, "P": 16, "Q": 55,
        },
        date_columns=("B", "C"),
    )
    _add_status_formatting(sessions_ws, "P", "Q", SESSION_COLORS)
    _add_table(sessions_ws, "SessionsTable")

    historical_ws = wb.create_sheet("Historical solved IDs")
    historical_ws.append(["Problem ID", "Imported", "Source"])
    for row in historical:
        historical_ws.append(
            [row["problem_id"], _excel_datetime(row["imported_at"]), row["source"]]
        )
    _style_sheet(
        historical_ws,
        {"A": 14, "B": 20, "C": 18},
        date_columns=("B",),
    )
    _add_table(historical_ws, "HistoricalSolvedTable")

    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = report_path.with_name(f".{report_path.stem}.tmp.xlsx")
    wb.save(temporary)
    os.replace(temporary, report_path)
    return report_path
